from io import BytesIO
from pathlib import Path
from struct import pack
from unittest.mock import patch
from zlib import compress, crc32

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.shared import Inches
from pypdf import PdfWriter

from app.services.document_validation import DocumentFormat
from app.services.parsing import (
    DocumentParseError,
    ParseFailureCode,
    ParserRegistry,
    TextExtractionStatus,
    create_default_parser_registry,
)
from app.services.parsing.docx_parser import DocxParser
from app.services.parsing.pdf_parser import PdfParser
from app.services.parsing.text_parser import MarkdownParser, PlainTextParser
from tests.sample_documents import make_image_backed_pdf, make_minimal_pdf


def make_blank_pdf() -> bytes:
    output = BytesIO()
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    writer.write(output)
    return output.getvalue()


def make_encrypted_pdf() -> bytes:
    output = BytesIO()
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    writer.encrypt("secret")
    writer.write(output)
    return output.getvalue()


def make_structured_docx() -> bytes:
    output = BytesIO()
    document = Document()
    document.add_heading("差旅费管理办法", level=1)
    document.add_heading("住宿费", level=2)
    document.add_paragraph("一线城市住宿上限为 600 元。")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "城市"
    table.cell(0, 1).text = "上限"
    table.cell(1, 0).text = "北京"
    table.cell(1, 1).text = "600 元"
    document.save(output)
    return output.getvalue()


def make_image_docx(tmp_path: Path) -> bytes:
    # 一像素 PNG；只用于确认图片边界，项目不会解释其视觉内容。
    def png_chunk(kind: bytes, data: bytes) -> bytes:
        return pack(">I", len(data)) + kind + data + pack(">I", crc32(kind + data))

    image_path = tmp_path / "pixel.png"
    image_path.write_bytes(
        b"\x89PNG\r\n\x1a\n"
        + png_chunk(b"IHDR", pack(">IIBBBBB", 1, 1, 8, 6, 0, 0, 0))
        + png_chunk(b"IDAT", compress(b"\x00\xff\xff\xff\xff"))
        + png_chunk(b"IEND", b"")
    )
    output = BytesIO()
    document = Document()
    document.add_paragraph("住宿费标准见正文。")
    document.add_picture(str(image_path), width=Inches(0.1))
    document.save(output)
    return output.getvalue()


def test_default_registry_parses_all_supported_formats() -> None:
    registry = create_default_parser_registry()

    pdf = registry.parse(DocumentFormat.PDF, BytesIO(make_minimal_pdf()))
    docx = registry.parse(DocumentFormat.DOCX, BytesIO(make_structured_docx()))
    markdown = registry.parse(DocumentFormat.MARKDOWN, BytesIO("# 住宿费\n上限 600 元".encode()))
    text = registry.parse(DocumentFormat.TEXT, BytesIO("报销制度正文".encode()))

    assert pdf.sections[0].page_number == 1
    assert docx.sections[-1].page_number is None
    assert docx.sections[-1].heading_path == ("差旅费管理办法", "住宿费")
    assert markdown.sections[0].heading_path == ("住宿费",)
    assert text.text_extraction_status is TextExtractionStatus.COMPLETE


def test_pdf_blank_page_has_no_text_but_does_not_claim_ocr() -> None:
    parsed = PdfParser().parse(BytesIO(make_blank_pdf()))

    assert parsed.text_extraction_status is TextExtractionStatus.NO_TEXT
    assert parsed.sections == ()
    assert parsed.needs_ocr is False
    assert parsed.warnings[0].code == "blank_page"


def test_pdf_image_only_page_is_no_text_and_requires_ocr() -> None:
    parsed = PdfParser().parse(BytesIO(make_image_backed_pdf((None,))))

    assert parsed.text_extraction_status is TextExtractionStatus.NO_TEXT
    assert parsed.image_only_page_numbers == (1,)
    assert parsed.needs_ocr is True
    assert parsed.sections == ()


def test_pdf_mixed_pages_only_indexes_extracted_text() -> None:
    parsed = PdfParser().parse(
        BytesIO(make_image_backed_pdf(("Hotel limit is 600 yuan for this policy.", None)))
    )

    assert parsed.text_extraction_status is TextExtractionStatus.PARTIAL
    assert [section.page_number for section in parsed.sections] == [1]
    assert parsed.image_only_page_numbers == (2,)
    assert parsed.needs_ocr is True


def test_pdf_with_decorative_image_and_enough_text_keeps_page() -> None:
    parsed = PdfParser().parse(
        BytesIO(make_image_backed_pdf(("Enough native policy text remains searchable.",)))
    )

    assert parsed.text_extraction_status is TextExtractionStatus.COMPLETE
    assert parsed.image_only_page_numbers == ()
    assert len(parsed.sections) == 1
    assert parsed.needs_ocr is False


def test_encrypted_and_corrupt_pdf_raise_safe_errors() -> None:
    with pytest.raises(DocumentParseError) as encrypted:
        PdfParser().parse(BytesIO(make_encrypted_pdf()))
    assert encrypted.value.code is ParseFailureCode.ENCRYPTED_DOCUMENT

    with pytest.raises(DocumentParseError) as corrupt:
        PdfParser().parse(BytesIO(b"%PDF-not-valid"))
    assert corrupt.value.code is ParseFailureCode.CORRUPT_DOCUMENT
    assert "not-valid" not in corrupt.value.safe_message


def test_docx_preserves_heading_path_and_table_without_fake_page() -> None:
    parsed = DocxParser().parse(BytesIO(make_structured_docx()))

    assert parsed.text_extraction_status is TextExtractionStatus.COMPLETE
    assert all(section.page_number is None for section in parsed.sections)
    assert any("北京 | 600 元" in section.text for section in parsed.sections)
    assert parsed.sections[-1].heading_path == ("差旅费管理办法", "住宿费")


def test_docx_with_image_warns_without_claiming_image_text(tmp_path: Path) -> None:
    parsed = DocxParser().parse(BytesIO(make_image_docx(tmp_path)))

    assert parsed.embedded_image_count == 1
    assert parsed.warnings[0].code == "embedded_images_detected"
    assert parsed.needs_ocr is False
    assert all("pixel" not in section.text for section in parsed.sections)


def test_markdown_images_are_not_fetched() -> None:
    content = b"# Policy\n![secret](file:///etc/passwd)\n![remote](https://example.invalid/x)\nText"
    with patch("pathlib.Path.read_bytes", side_effect=AssertionError("must not read image")):
        parsed = MarkdownParser().parse(BytesIO(content))

    assert parsed.embedded_image_count == 2
    assert parsed.warnings[0].code == "markdown_images_not_fetched"
    assert "file:///" not in parsed.sections[0].text
    assert "https://" not in parsed.sections[0].text


def test_plain_text_empty_content_has_no_sections() -> None:
    parsed = PlainTextParser().parse(BytesIO(b"  \n"))

    assert parsed.text_extraction_status is TextExtractionStatus.NO_TEXT
    assert parsed.sections == ()


def test_registry_rejects_duplicate_or_missing_parser() -> None:
    with pytest.raises(ValueError, match="already registered"):
        ParserRegistry((PlainTextParser(), PlainTextParser()))

    with pytest.raises(DocumentParseError) as missing:
        ParserRegistry().parse(DocumentFormat.TEXT, BytesIO(b"policy"))
    assert missing.value.code is ParseFailureCode.UNSUPPORTED_FORMAT


def test_docx_custom_heading_style_does_not_invent_page_number() -> None:
    output = BytesIO()
    document = Document()
    custom = document.styles.add_style("Heading 7 Custom", WD_STYLE_TYPE.PARAGRAPH)
    document.add_paragraph("附件要求", style=custom)
    document.add_paragraph("发票是必需附件。")
    document.save(output)

    parsed = DocxParser().parse(BytesIO(output.getvalue()))

    assert parsed.sections[-1].heading_path[-1] == "附件要求"
    assert parsed.sections[-1].page_number is None
